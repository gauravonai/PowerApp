"""
Build the JCB EDS Lab Material Portal canvas app sources.

    python3 tools/build.py            # writes app/Src/*.pa.yaml and paste/*

Outputs
  app/Src/                  pa.yaml source (Route A, packed into the .msapp by tools/pack.sh)
  paste/00-App.Formulas.txt paste into App > Formulas
  paste/01-App.OnStart.txt  paste into App > OnStart
  paste/02-App.StartScreen.txt
  paste/NN-<screen>.yaml    one "Paste code" per screen (Route B)
  tools/out/formulas.json   every formula, for the Power Fx syntax check
"""
import json
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

import pa  # noqa: E402
import appfx  # noqa: E402
import screens  # noqa: E402
import handtest  # noqa: E402

SRC = os.path.join(ROOT, "app", "Src")
PASTE = os.path.join(ROOT, "paste")
OUT = os.path.join(HERE, "out")


def sample_bom():
    """Three real parts with stock, one real part that is out, one part that does not exist."""
    try:
        import openpyxl
        wb = openpyxl.load_workbook(os.path.join(ROOT, "excel", "EDS-Lab-Data-PowerApps.xlsx"), data_only=False)
        parts = [r[0] for r in wb["Parts"].iter_rows(min_row=2, values_only=True) if r[0]]
        on = {}
        for r in wb["Movements"].iter_rows(min_row=2, values_only=True):
            if not r[1]:
                continue
            t, qn = str(r[2]).upper(), abs(float(r[3] or 0))
            s = 1 if t in ("RECEIPT", "RETURN", "ADJUST+") else -1 if t in ("ISSUE", "SCRAP", "ADJUST-") else 0
            on[str(r[1]).strip().upper()] = on.get(str(r[1]).strip().upper(), 0) + s * qn
        good = [p for p in parts if on.get(str(p).upper(), 0) >= 20 and " " not in str(p)][:3]
        out = [p for p in parts if on.get(str(p).upper(), 0) <= 0 and " " not in str(p)][:1]
        lines = ["%s\t%d" % (good[0], 10), "%s  %d" % (good[1], 5), "%s,%d" % (good[2], 2),
                 "%s\t%d" % (out[0], 4), "NEW/PART-0001\t3"]
    except Exception as e:  # workbook not there yet - fall back to a fixed sample
        print("sample BOM fallback:", e)
        lines = ["MC000RVT1CXX1SQ002\t10", "P_CON_GEN_ASLY_002  5", "NEW/PART-0001\t3"]
    return " & Char(10) & ".join(pa.q(l.replace("\t", "    ")) for l in lines)


def main():
    os.makedirs(SRC, exist_ok=True)
    os.makedirs(PASTE, exist_ok=True)
    os.makedirs(OUT, exist_ok=True)
    for d in (SRC, PASTE):
        for f in os.listdir(d):
            if f.endswith((".yaml", ".txt")):
                os.remove(os.path.join(d, f))

    formulas = appfx.formulas() + "\n\n// sample harness for BOM Compare (real part numbers from the workbook)\n" \
                                  "nfSampleBom = " + sample_bom() + ";"
    onstart = appfx.onstart()
    start = appfx.STARTSCREEN

    scr = screens.build_all()

    # ---------------------------------------------------------- name checks
    seen = {}
    for sname, _, root in scr:
        for c in root.walk():
            if c.name in seen:
                raise SystemExit("duplicate control name %s (in %s and %s)" % (c.name, seen[c.name], sname))
            seen[c.name] = sname
    names = set(seen) | {s for s, _, _ in scr}
    for must in ("lnRdC9", "lcC1", "lcC5", "lcC8", "tmC1", "tmC2", "tmC4", "stStkGal", "stInvGal"):
        if must not in names:
            raise SystemExit("formula references control %s that does not exist" % must)

    # ---------------------------------------------------------- Route A: pa.yaml
    with open(os.path.join(SRC, "App.pa.yaml"), "w") as f:
        f.write(pa.emit_app({"Formulas": formulas, "OnStart": onstart, "StartScreen": start}))
    with open(os.path.join(SRC, "_EditorState.pa.yaml"), "w") as f:
        f.write(pa.HEADER + "EditorState:\n  ScreensOrder:\n" + "".join("    - %s\n" % s for s, _, _ in scr))
    for sname, _, root in scr:
        with open(os.path.join(SRC, sname + ".pa.yaml"), "w") as f:
            f.write(pa.emit_screen(sname, {"Fill": "cBg" if sname != "scrPrint" else "RGBA(255, 255, 255, 1)"}, root))

    # ---------------------------------------------------------- Route B: paste files
    with open(os.path.join(PASTE, "00-App.Formulas.txt"), "w") as f:
        f.write(formulas + "\n")
    with open(os.path.join(PASTE, "01-App.OnStart.txt"), "w") as f:
        f.write(onstart + "\n")
    with open(os.path.join(PASTE, "02-App.StartScreen.txt"), "w") as f:
        f.write(start + "\n")
    for i, (sname, _, root) in enumerate(scr):
        with open(os.path.join(PASTE, "%02d-%s.yaml" % (i + 10, sname)), "w") as f:
            f.write(pa.emit_paste(root))

    # ---------------------------------------------------------- formulas for syntax check
    allf = [{"id": "App.Formulas", "f": formulas, "kind": "formulas"},
            {"id": "App.OnStart", "f": onstart}, {"id": "App.StartScreen", "f": start}]
    for sname, _, root in scr:
        for c in root.walk():
            for k, v in c.props.items():
                allf.append({"id": "%s.%s.%s" % (sname, c.name, k), "f": pa._fmt(v)})
    with open(os.path.join(OUT, "formulas.json"), "w") as f:
        json.dump(allf, f, ensure_ascii=False, indent=0)

    # ---------------------------------------------------------- flat model for the engine binding test
    model = {"formulas": formulas, "onstart": onstart, "startscreen": start, "screens": [s for s, _, _ in scr],
             "controls": [],
             "tests": {"role_canon": appfx.role_canon("RAWROLE"),
                       "hash_fx": appfx.hash_fx('"PW"', '"SALT"', "p1$"),
                       "hash_vectors": [[pw, salt, appfx.hash_py(pw, salt, "p1$")] for pw, salt in
                                        [("abc123", "gaurav.shelke"), ("Pune@2026", "akshay.aadarsh"),
                                         ('quote"and space 9', "x"), ("ünï 1a", "Mixed.Case")]],
                       "hand": {k: getattr(handtest, k) for k in ("PARTS", "MOVES", "ONHAND", "BOM")}}}

    def flat(c, sname, gal, parent):
        model["controls"].append({"name": c.name, "type": c.ctype, "screen": sname, "gallery": gal, "parent": parent,
                                  "props": {k: pa._fmt(v) for k, v in c.props.items()}})
        for ch in c.children:
            flat(ch, sname, c.name if c.ctype == pa.GALLERY else gal, c.name)
    for sname, _, root in scr:
        flat(root, sname, None, None)
    with open(os.path.join(OUT, "app.json"), "w") as f:
        json.dump(model, f, ensure_ascii=False)

    # ---------------------------------------------------------- Studio import guards (PA2105 / PA2108)
    from fix_src import check
    for d in (SRC, PASTE):
        for f in os.listdir(d):
            probs = check(open(os.path.join(d, f), encoding="utf-8").read())
            if probs:
                raise SystemExit("%s/%s would fail Studio import: %s" % (d, f, "; ".join(probs)))

    nctl = sum(1 for _, _, r in scr for _ in r.walk())
    print("screens: %d, controls: %d, formulas: %d" % (len(scr), nctl, len(allf)))
    print("App.Formulas: %d named formulas, %d chars" % (len(re.findall(r"^\s*\w+\s*=", formulas, re.M)), len(formulas)))


if __name__ == "__main__":
    main()
