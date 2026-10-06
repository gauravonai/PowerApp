"""v2 changes to the Power Apps workbook (safe to run more than once):
  1. Settings gets three Data Health rows. Excel formulas count the workbook's own rows, so the app
     can prove it read every row (About > Data Health and the sync badge in the header).
  2. The NewPartRequest SAMPLE row gets clearly non-numeric Plant Code / HSN text, so Power Apps types
     those columns as text (an HSN like 85366990 must be written as text).
  3. The POWER APPS note sheet explains multiple roles in one Role cell.

    python3 tools/upgrade_excel_v2.py [excel/EDS-Lab-Data-PowerApps.xlsx]
"""
import sys

import openpyxl

HEALTH = [
    ("ExcelPartRows", "=ROWS(tblParts)", "Data Health: rows in Parts, counted by Excel. Do not edit."),
    ("ExcelMovementRows", "=ROWS(tblMoves)", "Data Health: rows in Movements, counted by Excel. Do not edit."),
    ("ExcelStockTotal",
     '=SUMPRODUCT((tblMoves[Type]="RECEIPT")+(tblMoves[Type]="RETURN")+(tblMoves[Type]="ADJUST+")'
     '-(tblMoves[Type]="ISSUE")-(tblMoves[Type]="SCRAP")-(tblMoves[Type]="ADJUST-"),ABS(tblMoves[Qty]))',
     "Data Health: total stock of all parts = signed sum of Movements. Do not edit."),
]


def table_of(wb, name):
    for ws in wb.worksheets:
        if name in ws.tables:
            return ws, ws.tables[name]
    raise KeyError(name)


def main(path):
    wb = openpyxl.load_workbook(path)
    ws, t = table_of(wb, "tblSettings")
    a, b = t.ref.split(":")
    last = int("".join(ch for ch in b if ch.isdigit()))
    have = {ws.cell(row=r, column=1).value for r in range(2, last + 1)}
    for key, f, note in HEALTH:
        if key in have:
            continue
        last += 1
        ws.cell(row=last, column=1, value=key)
        ws.cell(row=last, column=2, value=f)
        ws.cell(row=last, column=3, value=note)
    t.ref = "%s:%s%d" % (a, "".join(ch for ch in b if ch.isalpha()), last)
    if t.autoFilter is not None:
        t.autoFilter.ref = t.ref

    ws, t = table_of(wb, "tblNewPart")
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(row=r, column=2).value or "").startswith("SAMPLE"):
            ws.cell(row=r, column=6, value="SAMPLE-PLANT")
            ws.cell(row=r, column=13, value="SAMPLE-HSN")

    n = wb["POWER APPS"]
    if not any(str(c.value or "").startswith("ROLES:") for c in n["B"]):
        row = n.max_row + 2
        for i, l in enumerate([
            "ROLES: one person can hold several roles. Write them in the Role cell separated by commas, for example",
            "'Engineer, Manager'. The first one is the usual role. Lab Lead and Lab Admin can view every role.",
            "",
            "DATA HEALTH: the three Excel* rows at the bottom of Settings are formulas. Leave them; the app compares them",
            "with what it read (About > Data Health)."]):
            n.cell(row=row + i, column=2, value=l)
    wb.save(path)
    print("upgraded", path)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "excel/EDS-Lab-Data-PowerApps.xlsx")
